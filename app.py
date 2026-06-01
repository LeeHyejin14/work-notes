import tkinter as tk
from tkinter import ttk, messagebox, filedialog, colorchooser
import json as _json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime
from PIL import Image, ImageTk, ImageGrab
import os
import sys
import shutil
import tempfile

if getattr(sys, "frozen", False):
    BASE_DIR = os.path.join(os.path.expanduser("~"), "Documents", "업무노트")
else:
    BASE_DIR = os.path.dirname(os.path.abspath(__file__))
EXCEL_PATH = os.path.join(BASE_DIR, "업무교육.xlsx")
IMG_DIR    = os.path.join(BASE_DIR, "images")
COLUMNS    = ["ID", "날짜", "카테고리", "태그", "제목", "내용", "이미지"]
THUMB_SIZE = (120, 90)

os.makedirs(IMG_DIR, exist_ok=True)


# ── 이미지 유틸 ────────────────────────────────────────────────────────────

def images_for(rec):
    """레코드의 이미지 파일명 리스트 반환."""
    raw = rec.get("이미지", "").strip()
    return [f for f in raw.split(",") if f.strip()] if raw else []

def save_image(src_path, record_id):
    """이미지를 images/{id}/ 폴더에 복사하고 파일명 반환."""
    folder = os.path.join(IMG_DIR, str(record_id))
    os.makedirs(folder, exist_ok=True)
    fname = os.path.basename(src_path)
    # 중복 파일명 처리
    dst = os.path.join(folder, fname)
    if os.path.exists(dst):
        base, ext = os.path.splitext(fname)
        fname = f"{base}_{int(datetime.now().timestamp())}{ext}"
        dst = os.path.join(folder, fname)
    shutil.copy2(src_path, dst)
    return fname

def full_image_path(record_id, fname):
    return os.path.join(IMG_DIR, str(record_id), fname)

def make_thumbnail(path):
    """PIL 썸네일 ImageTk 반환. 실패 시 None."""
    try:
        img = Image.open(path)
        img.thumbnail(THUMB_SIZE, Image.LANCZOS)
        return ImageTk.PhotoImage(img)
    except Exception:
        return None

# ── 서식 텍스트 직렬화 ────────────────────────────────────────────────────

_FMTMARK = "\x00FMT\x00"

def _setup_fmt_tags(tw, size=10):
    fam = "Arial"
    tw.tag_configure("fmt_b",  font=(fam, size, "bold"))
    tw.tag_configure("fmt_i",  font=(fam, size, "italic"))
    tw.tag_configure("fmt_bi", font=(fam, size, "bold italic"))
    tw.tag_configure("fmt_u",  underline=True)

def _ensure_fmt_tag(tw, name, size=10):
    if name in tw.tag_names():
        return
    fam = "Arial"
    if   name == "fmt_b":              tw.tag_configure(name, font=(fam, size, "bold"))
    elif name == "fmt_i":              tw.tag_configure(name, font=(fam, size, "italic"))
    elif name == "fmt_bi":             tw.tag_configure(name, font=(fam, size, "bold italic"))
    elif name == "fmt_u":              tw.tag_configure(name, underline=True)
    elif name.startswith("fmt_c_"):    tw.tag_configure(name, foreground=name[6:])
    elif name.startswith("fmt_s_"):    tw.tag_configure(name, font=(fam, int(name[6:])))

def _content_preview(stored, n=50):
    if stored and stored.startswith(_FMTMARK):
        try:
            text = _json.loads(stored[len(_FMTMARK):])["t"]
        except Exception:
            text = ""
    else:
        text = stored or ""
    return text.replace("\n", " ")[:n]

def content_serialize(tw):
    """Text 위젯 내용 + 서식을 저장용 문자열로 변환."""
    text = tw.get("1.0", "end-1c")
    tags = []
    for name in tw.tag_names():
        if not name.startswith("fmt_"):
            continue
        ranges = tw.tag_ranges(name)
        if not ranges:
            continue
        pairs = []
        for j in range(0, len(ranges), 2):
            r0 = tw.count("1.0", ranges[j],   "chars")
            r1 = tw.count("1.0", ranges[j+1], "chars")
            s = (r0[0] if isinstance(r0, tuple) else r0) or 0
            e = (r1[0] if isinstance(r1, tuple) else r1) or 0
            pairs.append([s, e])
        tags.append([name, pairs])
    if not tags:
        return text
    return _FMTMARK + _json.dumps({"t": text, "f": tags}, ensure_ascii=False)

def content_load(tw, stored, size=10):
    """저장된 문자열을 Text 위젯에 로드 (서식 포함)."""
    tw.delete("1.0", "end")
    _setup_fmt_tags(tw, size)
    if not stored or not stored.startswith(_FMTMARK):
        tw.insert("1.0", stored or "")
        return
    data = _json.loads(stored[len(_FMTMARK):])
    tw.insert("1.0", data["t"])
    for name, pairs in data["f"]:
        _ensure_fmt_tag(tw, name, size)
        for s, e in pairs:
            tw.tag_add(name, f"1.0 + {s} chars", f"1.0 + {e} chars")


# ── 항목 추가/수정 다이얼로그 ──────────────────────────────────────────────

class EntryDialog(tk.Toplevel):
    def __init__(self, parent, title="항목 추가", data=None, record_id=None):
        super().__init__(parent)
        self.title(title)
        self.resizable(True, True)
        self.result = None
        self.record_id = record_id  # 수정 시 기존 ID, 추가 시 None
        self._pending_images = []   # (src_path,) 새로 추가된 이미지
        self._existing_images = list(images_for(data or {}))  # 기존 이미지 파일명
        self._removed_images = []
        self._thumb_refs = []       # GC 방지
        self._build(data or {})
        self.transient(parent)
        self.grab_set()
        self.geometry("560x560")
        self.minsize(460, 420)
        self.protocol("WM_DELETE_WINDOW", self._cancel)
        self.after(50, self._center)

    def _center(self):
        self.update_idletasks()
        x = self.master.winfo_x() + (self.master.winfo_width()  - self.winfo_width())  // 2
        y = self.master.winfo_y() + (self.master.winfo_height() - self.winfo_height()) // 2
        self.geometry(f"+{x}+{y}")

    def _build(self, data):
        pad = {"padx": 12, "pady": 4}
        frame = ttk.Frame(self, padding=12)
        frame.pack(fill="both", expand=True)
        frame.columnconfigure(1, weight=1)

        # 텍스트 필드
        fields = [("날짜", "날짜"), ("카테고리", "카테고리"), ("태그", "태그"), ("제목", "제목")]
        self.vars = {}
        for i, (label, key) in enumerate(fields):
            ttk.Label(frame, text=label).grid(row=i, column=0, sticky="nw", **pad)
            default = data.get(key, "")
            if key == "날짜" and not default:
                default = datetime.now().strftime("%Y-%m-%d")
            var = tk.StringVar(value=default)
            row_frame = ttk.Frame(frame)
            row_frame.grid(row=i, column=1, sticky="ew", **pad)
            row_frame.columnconfigure(0, weight=1)
            entry = ttk.Entry(row_frame, textvariable=var, font=("Arial", 10))
            entry.grid(row=0, column=0, sticky="ew")
            if key == "날짜":
                ttk.Label(row_frame, text="(YYYY-MM-DD)", foreground="#888",
                          font=("Arial", 9)).grid(row=0, column=1, padx=(6, 0))
                entry.focus_set()
            self.vars[key] = var
            if key == "카테고리":
                entry.focus_set()

        # 내용
        ttk.Label(frame, text="내용").grid(row=4, column=0, sticky="nw", **pad)
        cw = ttk.Frame(frame)
        cw.grid(row=4, column=1, sticky="nsew", **pad)
        cw.rowconfigure(1, weight=1)
        cw.columnconfigure(0, weight=1)
        frame.rowconfigure(4, weight=1)

        # 서식 툴바
        fb = ttk.Frame(cw)
        fb.grid(row=0, column=0, columnspan=2, sticky="w", pady=(0, 2))
        tk.Button(fb, text="B", font=("Arial", 10, "bold"),   width=2, relief="groove",
                  command=lambda: self._fmt_toggle("fmt_b")).pack(side="left", padx=1)
        tk.Button(fb, text="I", font=("Arial", 10, "italic"), width=2, relief="groove",
                  command=lambda: self._fmt_toggle("fmt_i")).pack(side="left", padx=1)
        tk.Button(fb, text="U", font=("Arial", 10),           width=2, relief="groove",
                  command=lambda: self._fmt_toggle("fmt_u")).pack(side="left", padx=1)
        ttk.Separator(fb, orient="vertical").pack(side="left", fill="y", padx=5)
        self._size_var = tk.StringVar(value="10")
        sz = ttk.Combobox(fb, textvariable=self._size_var, state="readonly", width=4,
                          values=["8","9","10","11","12","14","16","18","20","24"])
        sz.pack(side="left", padx=2)
        sz.bind("<<ComboboxSelected>>", lambda _: self._fmt_size(int(self._size_var.get())))
        ttk.Separator(fb, orient="vertical").pack(side="left", fill="y", padx=5)
        tk.Button(fb, text="색상", font=("Arial", 9), relief="groove",
                  command=self._fmt_color).pack(side="left", padx=1)
        tk.Button(fb, text="초기화", font=("Arial", 9), relief="groove",
                  command=self._fmt_clear).pack(side="left", padx=4)

        self.content_text = tk.Text(cw, height=6, wrap="word", font=("Arial", 10))
        ct_sb = ttk.Scrollbar(cw, orient="vertical", command=self.content_text.yview)
        self.content_text.configure(yscrollcommand=ct_sb.set)
        self.content_text.grid(row=1, column=0, sticky="nsew")
        ct_sb.grid(row=1, column=1, sticky="ns")
        _setup_fmt_tags(self.content_text)
        content_load(self.content_text, data.get("내용", ""))

        tw = self.content_text
        tw.bind("<Control-b>", lambda _: self._fmt_toggle("fmt_b") or "break")
        tw.bind("<Control-i>", lambda _: self._fmt_toggle("fmt_i") or "break")
        tw.bind("<Control-u>", lambda _: self._fmt_toggle("fmt_u") or "break")

        # 이미지 섹션
        img_label_frame = ttk.LabelFrame(frame, text="첨부 이미지", padding=6)
        img_label_frame.grid(row=5, column=0, columnspan=2, sticky="ew", padx=12, pady=6)
        img_label_frame.columnconfigure(0, weight=1)
        frame.rowconfigure(5, weight=0)

        # 이미지 추가 버튼
        btn_row = ttk.Frame(img_label_frame)
        btn_row.grid(row=0, column=0, sticky="w")
        ttk.Button(btn_row, text="+ 이미지 추가", command=self._add_image).pack(side="left", padx=2)
        ttk.Button(btn_row, text="📋 클립보드 붙여넣기", command=self._paste_clipboard).pack(side="left", padx=2)

        # 썸네일 캔버스
        self.thumb_frame = ttk.Frame(img_label_frame)
        self.thumb_frame.grid(row=1, column=0, sticky="ew", pady=4)
        self._render_thumbs()

        # 버튼
        btn_frame = ttk.Frame(self)
        btn_frame.pack(fill="x", padx=12, pady=8)
        ttk.Button(btn_frame, text="저장", command=self._save).pack(side="right", padx=4)
        ttk.Button(btn_frame, text="취소", command=self._cancel).pack(side="right", padx=4)

    def _render_thumbs(self):
        for w in self.thumb_frame.winfo_children():
            w.destroy()
        self._thumb_refs.clear()

        all_items = []
        # 기존 이미지
        for fname in self._existing_images:
            all_items.append(("existing", fname))
        # 새로 추가
        for src in self._pending_images:
            all_items.append(("pending", src))

        for col, (kind, val) in enumerate(all_items):
            cell = ttk.Frame(self.thumb_frame)
            cell.grid(row=0, column=col, padx=4, pady=2)

            if kind == "existing" and self.record_id:
                path = full_image_path(self.record_id, val)
                name = val
            else:
                path = val
                name = os.path.basename(val)

            thumb = make_thumbnail(path)
            if thumb:
                self._thumb_refs.append(thumb)
                lbl = tk.Label(cell, image=thumb, cursor="hand2")
                lbl.pack()
                lbl.bind("<Button-1>", lambda e, p=path: self._show_image_inline(p))
            else:
                tk.Label(cell, text="[이미지\n없음]", width=10, height=4,
                         relief="groove").pack()

            name_short = name if len(name) <= 14 else name[:12] + "…"
            tk.Label(cell, text=name_short, font=("Arial", 9)).pack()

            # 삭제 버튼
            tk.Button(cell, text="✕", font=("Arial", 8), fg="red", bd=0,
                      command=lambda k=kind, v=val: self._remove_image(k, v)).pack()

        if not all_items:
            ttk.Label(self.thumb_frame, text="이미지 없음", foreground="#aaa").grid(row=0, column=0)

    def _add_image(self):
        paths = filedialog.askopenfilenames(
            title="이미지 선택",
            filetypes=[("이미지 파일", "*.png *.jpg *.jpeg *.gif *.bmp *.webp"), ("모든 파일", "*.*")]
        )
        for p in paths:
            self._pending_images.append(p)
        self._render_thumbs()

    def _paste_clipboard(self):
        if isinstance(self.focus_get(), (tk.Text, tk.Entry, ttk.Entry)):
            return
        try:
            img = ImageGrab.grabclipboard()
        except Exception as e:
            messagebox.showerror("오류", f"클립보드 접근 실패:\n{e}", parent=self)
            return
        if isinstance(img, list):
            # 파일 경로 목록이 복사된 경우 이미지 파일만 추가
            added = False
            for p in img:
                if p.lower().endswith((".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp")):
                    self._pending_images.append(p)
                    added = True
            if not added:
                return
        elif isinstance(img, Image.Image):
            tmp = tempfile.NamedTemporaryFile(
                suffix=".png", prefix="clip_", delete=False,
                dir=tempfile.gettempdir()
            )
            img.save(tmp.name, "PNG")
            tmp.close()
            self._pending_images.append(tmp.name)
        else:
            return
        self._render_thumbs()

    def _remove_image(self, kind, val):
        if kind == "existing":
            self._existing_images.remove(val)
            self._removed_images.append(val)
        else:
            self._pending_images.remove(val)
        self._render_thumbs()

    def _fmt_toggle(self, tag):
        tw = self.content_text
        try:
            s, e = tw.index("sel.first"), tw.index("sel.last")
        except tk.TclError:
            return
        if tag in ("fmt_b", "fmt_i"):
            other = "fmt_i" if tag == "fmt_b" else "fmt_b"
            cur = set(tw.tag_names(s))
            if "fmt_bi" in cur:
                tw.tag_remove("fmt_bi", s, e)
                tw.tag_add(other, s, e)
            elif tag in cur:
                tw.tag_remove(tag, s, e)
            elif other in cur:
                tw.tag_remove(other, s, e)
                tw.tag_add("fmt_bi", s, e)
            else:
                tw.tag_add(tag, s, e)
        else:
            if tag in set(tw.tag_names(s)):
                tw.tag_remove(tag, s, e)
            else:
                tw.tag_add(tag, s, e)

    def _fmt_color(self):
        tw = self.content_text
        try:
            s, e = tw.index("sel.first"), tw.index("sel.last")
        except tk.TclError:
            return
        color = colorchooser.askcolor(parent=self, title="글자 색 선택")[1]
        if not color:
            return
        name = f"fmt_c_{color}"
        _ensure_fmt_tag(tw, name)
        tw.tag_add(name, s, e)

    def _fmt_size(self, size):
        tw = self.content_text
        try:
            s, e = tw.index("sel.first"), tw.index("sel.last")
        except tk.TclError:
            return
        for t in tw.tag_names():
            if t.startswith("fmt_s_"):
                tw.tag_remove(t, s, e)
        name = f"fmt_s_{size}"
        _ensure_fmt_tag(tw, name)
        tw.tag_add(name, s, e)

    def _fmt_clear(self):
        tw = self.content_text
        try:
            s, e = tw.index("sel.first"), tw.index("sel.last")
        except tk.TclError:
            s, e = "1.0", "end"
        for t in list(tw.tag_names()):
            if t.startswith("fmt_"):
                tw.tag_remove(t, s, e)

    def _show_image_inline(self, path):
        for w in self.thumb_frame.winfo_children():
            w.destroy()
        self._thumb_refs.clear()

        ttk.Button(self.thumb_frame, text="← 목록",
                   command=self._render_thumbs).grid(row=0, column=0, sticky="w", pady=(2, 4))
        try:
            img = Image.open(path)
            self.update_idletasks()
            max_w = max(self.thumb_frame.winfo_width() - 10, 200)
            max_h = 300
            img.thumbnail((max_w, max_h), Image.LANCZOS)
            photo = ImageTk.PhotoImage(img)
            self._thumb_refs.append(photo)
            tk.Label(self.thumb_frame, image=photo).grid(row=1, column=0, padx=4, pady=4)
        except Exception as e:
            ttk.Label(self.thumb_frame, text=f"로드 실패: {e}").grid(row=1, column=0, pady=10)

    def _save(self):
        if not self.vars["제목"].get().strip():
            messagebox.showwarning("입력 오류", "제목은 필수입니다.", parent=self)
            return
        date_str = self.vars["날짜"].get().strip()
        try:
            datetime.strptime(date_str, "%Y-%m-%d")
        except ValueError:
            messagebox.showwarning("입력 오류", "날짜 형식이 올바르지 않습니다.\nYYYY-MM-DD 형식으로 입력해 주세요.", parent=self)
            return
        self.result = {
            "날짜":          date_str,
            "카테고리":      self.vars["카테고리"].get().strip(),
            "태그":          self.vars["태그"].get().strip(),
            "제목":          self.vars["제목"].get().strip(),
            "내용":          content_serialize(self.content_text),
            "_existing":     list(self._existing_images),
            "_pending":      list(self._pending_images),
            "_removed":      list(self._removed_images),
        }
        self.destroy()

    def _cancel(self):
        self.destroy()


# ── 메인 앱 ───────────────────────────────────────────────────────────────

class WorkNotesApp:
    def __init__(self, root):
        self.root = root
        self.root.title("업무교육")
        self.root.geometry("1150x700")
        self.root.minsize(800, 520)

        self.records  = []
        self.filtered = []
        self.next_id  = 1
        self._thumb_refs = []  # GC 방지

        self._build_ui()
        self._load_excel()
        self._refresh_categories()
        self._apply_filter()

    # ── UI ────────────────────────────────────────────────────────────────

    def _build_ui(self):
        # 툴바
        tb = ttk.Frame(self.root, padding=(8, 6))
        tb.pack(fill="x", side="top")

        ttk.Button(tb, text="+ 추가",    command=self._add).pack(side="left", padx=2)
        ttk.Button(tb, text="✏ 수정",    command=self._edit).pack(side="left", padx=2)
        ttk.Button(tb, text="🗑 삭제",   command=self._delete).pack(side="left", padx=2)
        ttk.Separator(tb, orient="vertical").pack(side="left", fill="y", padx=6)
        ttk.Button(tb, text="↑ 파일 가져오기", command=self._import_excel).pack(side="left", padx=2)
        ttk.Button(tb, text="↓ 파일 내보내기", command=self._export_excel).pack(side="left", padx=2)

        sf = ttk.Frame(tb)
        sf.pack(side="right")
        ttk.Label(sf, text="검색:").pack(side="left")
        self.search_var = tk.StringVar()
        self.search_var.trace_add("write", lambda *_: self._apply_filter())
        ttk.Entry(sf, textvariable=self.search_var, width=22).pack(side="left", padx=4)
        ttk.Button(sf, text="✕", width=2,
                   command=lambda: self.search_var.set("")).pack(side="left")

        # 메인 패널
        paned = ttk.PanedWindow(self.root, orient="horizontal")
        paned.pack(fill="both", expand=True, padx=8, pady=(0, 4))

        # 왼쪽: 카테고리
        left = ttk.Frame(paned, width=150)
        paned.add(left, weight=0)
        ttk.Label(left, text="카테고리", font=("Arial", 11, "bold")).pack(anchor="w", padx=6, pady=(6, 2))
        self.cat_list = tk.Listbox(left, selectmode="single", activestyle="none",
                                   font=("Arial", 11), bd=0, highlightthickness=1)
        self.cat_list.pack(fill="both", expand=True, padx=4, pady=4)
        self.cat_list.bind("<<ListboxSelect>>", lambda _: self._apply_filter())

        # 오른쪽: 테이블 + 상세
        right = ttk.Frame(paned)
        paned.add(right, weight=1)
        right.rowconfigure(0, weight=2)
        right.rowconfigure(1, weight=1)
        right.columnconfigure(0, weight=1)

        # 테이블
        tree_frame = ttk.Frame(right)
        tree_frame.grid(row=0, column=0, sticky="nsew")
        tree_frame.rowconfigure(0, weight=1)
        tree_frame.columnconfigure(0, weight=1)

        cols = ("날짜", "카테고리", "태그", "제목", "내용", "📎")
        self.tree = ttk.Treeview(tree_frame, columns=cols, show="headings", selectmode="browse")
        widths = [90, 100, 100, 150, 220, 40]
        for col, w in zip(cols, widths):
            self.tree.heading(col, text=col, command=lambda c=col: self._sort(c))
            self.tree.column(col, width=w, minwidth=40)

        vsb = ttk.Scrollbar(tree_frame, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=vsb.set)
        self.tree.grid(row=0, column=0, sticky="nsew")
        vsb.grid(row=0, column=1, sticky="ns")

        self.tree.bind("<<TreeviewSelect>>", self._on_select)
        self.tree.bind("<Double-1>", lambda _: self._edit())

        # 상세 패널 (내용 + 이미지)
        detail_paned = ttk.PanedWindow(right, orient="horizontal")
        detail_paned.grid(row=1, column=0, sticky="nsew", pady=(6, 0))

        # 내용 텍스트
        content_frame = ttk.LabelFrame(detail_paned, text="내용", padding=8)
        detail_paned.add(content_frame, weight=2)
        content_frame.rowconfigure(0, weight=1)
        content_frame.columnconfigure(0, weight=1)

        self.detail_text = tk.Text(content_frame, wrap="word", font=("Arial", 10),
                                   state="disabled", bd=0, bg="#F9F9F9")
        _setup_fmt_tags(self.detail_text)
        dsb = ttk.Scrollbar(content_frame, orient="vertical", command=self.detail_text.yview)
        self.detail_text.configure(yscrollcommand=dsb.set)
        self.detail_text.grid(row=0, column=0, sticky="nsew")
        dsb.grid(row=0, column=1, sticky="ns")

        # 이미지 패널
        img_outer = ttk.LabelFrame(detail_paned, text="이미지", padding=8)
        detail_paned.add(img_outer, weight=1)
        img_outer.rowconfigure(0, weight=1)
        img_outer.columnconfigure(0, weight=1)

        self.img_canvas = tk.Canvas(img_outer, bg="#F9F9F9", bd=0, highlightthickness=0)
        img_vsb = ttk.Scrollbar(img_outer, orient="vertical", command=self.img_canvas.yview)
        self.img_canvas.configure(yscrollcommand=img_vsb.set)
        self.img_canvas.grid(row=0, column=0, sticky="nsew")
        img_vsb.grid(row=0, column=1, sticky="ns")

        self.img_inner = ttk.Frame(self.img_canvas)
        self.img_canvas_window = self.img_canvas.create_window((0, 0), window=self.img_inner, anchor="nw")
        self.img_inner.bind("<Configure>", self._on_img_frame_configure)
        self.img_canvas.bind("<Configure>", self._on_img_canvas_configure)

        # 상태바
        self.status_var = tk.StringVar()
        ttk.Label(self.root, textvariable=self.status_var, anchor="w",
                  font=("Arial", 10), foreground="#555").pack(fill="x", padx=8, pady=(0, 4))

    def _on_img_frame_configure(self, _e):
        self.img_canvas.configure(scrollregion=self.img_canvas.bbox("all"))

    def _on_img_canvas_configure(self, e):
        self.img_canvas.itemconfig(self.img_canvas_window, width=e.width)

    # ── 이미지 패널 렌더링 ────────────────────────────────────────────────

    def _show_images(self, rec):
        for w in self.img_inner.winfo_children():
            w.destroy()
        self._thumb_refs.clear()

        if rec is None:
            return

        fnames = images_for(rec)
        if not fnames:
            ttk.Label(self.img_inner, text="이미지 없음", foreground="#aaa",
                      font=("Arial", 10)).pack(pady=20)
            return

        for fname in fnames:
            path = full_image_path(rec["ID"], fname)
            if not os.path.exists(path):
                continue
            cell = ttk.Frame(self.img_inner)
            cell.pack(fill="x", pady=4, padx=4)

            thumb = make_thumbnail(path)
            if thumb:
                self._thumb_refs.append(thumb)
                lbl = tk.Label(cell, image=thumb, cursor="hand2")
                lbl.pack()
                lbl.bind("<Button-1>", lambda e, p=path, r=rec: self._show_full_image(p, r))
            fname_short = fname if len(fname) <= 20 else fname[:18] + "…"
            tk.Label(cell, text=fname_short, font=("Arial", 9),
                     foreground="#555").pack()

    def _show_full_image(self, path, rec):
        for w in self.img_inner.winfo_children():
            w.destroy()
        self._thumb_refs.clear()

        ttk.Button(self.img_inner, text="← 목록",
                   command=lambda: self._show_images(rec)).pack(anchor="w", padx=4, pady=(4, 2))
        try:
            img = Image.open(path)
            self.img_canvas.update_idletasks()
            panel_w = max(self.img_canvas.winfo_width() - 20, 100)
            if img.width > panel_w:
                img = img.resize(
                    (panel_w, int(img.height * panel_w / img.width)), Image.LANCZOS
                )
            photo = ImageTk.PhotoImage(img)
            self._thumb_refs.append(photo)
            tk.Label(self.img_inner, image=photo).pack(padx=4, pady=4)
            fname = os.path.basename(path)
            ttk.Label(self.img_inner,
                      text=fname if len(fname) <= 28 else fname[:26] + "…",
                      foreground="#555", font=("Arial", 9)).pack()
        except Exception as e:
            ttk.Label(self.img_inner, text=f"로드 실패: {e}").pack(pady=10)

    # ── 데이터 로드/저장 ──────────────────────────────────────────────────

    def _seed_sample_data(self):
        samples = [
            ("2026-05-01", "교육", "파이썬,자동화", "파이썬 기초 교육 수강",
             "파이썬 기초 문법 및 자동화 스크립트 작성 방법 교육 수강.\n"
             "- 변수, 조건문, 반복문 복습\n- 파일 입출력 실습\n- 업무 자동화 예제 3가지 실습"),
            ("2026-05-05", "회의", "주간회의,기획", "5월 1주차 팀 주간 회의",
             "참석: 팀장, 팀원 5명\n\n주요 안건:\n1. 4월 실적 공유 — 목표 대비 92% 달성\n"
             "2. 5월 일정 조율 — 중간 점검일 5/15로 확정\n3. 신규 프로젝트 역할 분담 논의"),
            ("2026-05-08", "업무", "엑셀,보고서", "월간 업무 보고서 작성",
             "5월 월간 보고서 초안 작성 완료.\n"
             "- 실적 데이터 정리 (엑셀)\n- 부서별 KPI 취합\n- 팀장 검토 후 수정 예정"),
            ("2026-05-12", "교육", "보안,필수교육", "개인정보 보호 의무 교육",
             "연 1회 의무 교육 이수.\n\n주요 내용:\n- 개인정보 처리 방침 변경 사항\n"
             "- 유출 사고 사례 및 예방 방법\n- 위반 시 처벌 기준\n\n이수증 첨부 완료"),
            ("2026-05-14", "업무", "고객,미팅", "A사 미팅 결과 정리",
             "미팅 일시: 2026-05-14 14:00\n장소: A사 회의실\n\n논의 사항:\n"
             "- 계약 갱신 조건 협의 (단가 조정 요청)\n- 추가 요구 사항: 월별 리포트 제공\n"
             "- 다음 미팅 6/3으로 예약"),
            ("2026-05-16", "회의", "부서,협업", "타 부서 협업 킥오프 미팅",
             "마케팅팀 + 개발팀 협업 프로젝트 시작 회의.\n\n결정 사항:\n"
             "- 담당자: 마케팅 2명, 개발 3명\n- 주간 싱크 매주 화요일 10시\n- 1차 산출물 목표: 5/30"),
            ("2026-05-19", "업무", "인수인계,정리", "업무 노트 프로그램 사용 시작",
             "업무 노트 프로그램 도입 및 초기 세팅 완료.\n\n활용 계획:\n"
             "- 교육 이수 내용 기록\n- 회의록 보관\n- 업무 진행 사항 추적\n- 이미지/캡처 첨부 활용"),
        ]
        for i, (date, cat, tag, title, content) in enumerate(samples, 1):
            self.records.append({
                "ID": str(i), "날짜": date, "카테고리": cat,
                "태그": tag, "제목": title, "내용": content, "이미지": "",
            })
        self.next_id = len(samples) + 1
        self._save_excel(silent=True)
        self._set_status(f"샘플 데이터 {len(samples)}개 추가됨")

    def _load_excel(self):
        if not os.path.exists(EXCEL_PATH):
            self._seed_sample_data()
            return
        try:
            wb = openpyxl.load_workbook(EXCEL_PATH)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                return
            header = [str(c) for c in rows[0]]
            max_id = 0
            for row in rows[1:]:
                if not any(row):
                    continue
                d = dict(zip(header, [str(v) if v is not None else "" for v in row]))
                try:
                    max_id = max(max_id, int(d.get("ID", 0)))
                except ValueError:
                    pass
                self.records.append(d)
            self.next_id = max_id + 1
            self._set_status(f"파일 로드 완료 — {len(self.records)}개 항목")
        except Exception as e:
            messagebox.showerror("오류", f"파일을 불러오지 못했습니다:\n{e}")

    def _save_excel(self, silent=False):
        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "업무교육"
            header_fill = PatternFill("solid", fgColor="4A90D9")
            header_font = Font(bold=True, color="FFFFFF")
            col_w = [8, 14, 16, 20, 36, 60, 30]
            for ci, (col, w) in enumerate(zip(COLUMNS, col_w), 1):
                cell = ws.cell(row=1, column=ci, value=col)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center")
                ws.column_dimensions[cell.column_letter].width = w
            for ri, rec in enumerate(self.records, 2):
                for ci, col in enumerate(COLUMNS, 1):
                    ws.cell(row=ri, column=ci, value=rec.get(col, ""))
            wb.save(EXCEL_PATH)
            if not silent:
                self._set_status(f"저장 완료 ({EXCEL_PATH})")
        except Exception as e:
            messagebox.showerror("저장 오류", str(e))

    def _export_excel(self):
        if not self.filtered:
            messagebox.showinfo("알림", "내보낼 항목이 없습니다.")
            return
        path = filedialog.asksaveasfilename(
            title="엑셀로 내보내기",
            defaultextension=".xlsx",
            initialfile=f"업무교육_내보내기_{datetime.now().strftime('%Y%m%d')}.xlsx",
            filetypes=[("Excel 파일", "*.xlsx"), ("모든 파일", "*.*")],
        )
        if not path:
            return
        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "업무교육"
            header_fill = PatternFill("solid", fgColor="4A90D9")
            header_font = Font(bold=True, color="FFFFFF")
            export_cols = ["날짜", "카테고리", "태그", "제목", "내용"]
            col_widths   = [14,     16,         20,    36,    60]
            for ci, (col, w) in enumerate(zip(export_cols, col_widths), 1):
                cell = ws.cell(row=1, column=ci, value=col)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center")
                ws.column_dimensions[cell.column_letter].width = w
            for ri, rec in enumerate(self.filtered, 2):
                for ci, col in enumerate(export_cols, 1):
                    c = ws.cell(row=ri, column=ci, value=rec.get(col, ""))
                    c.alignment = Alignment(wrap_text=True, vertical="top")
            wb.save(path)
            self._set_status(f"{len(self.filtered)}개 항목 내보내기 완료 → {os.path.basename(path)}")
        except Exception as e:
            messagebox.showerror("내보내기 오류", str(e))

    def _import_excel(self):
        path = filedialog.askopenfilename(
            title="엑셀 파일 선택",
            filetypes=[("Excel 파일", "*.xlsx *.xls"), ("모든 파일", "*.*")]
        )
        if not path:
            return
        try:
            wb = openpyxl.load_workbook(path)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                return
            header = [str(c) for c in rows[0]]
            imported = 0
            for row in rows[1:]:
                if not any(row):
                    continue
                d = dict(zip(header, [str(v) if v is not None else "" for v in row]))
                d["ID"] = str(self.next_id)
                self.next_id += 1
                if not d.get("날짜"):
                    d["날짜"] = datetime.now().strftime("%Y-%m-%d")
                self.records.append(d)
                imported += 1
            self._refresh_categories()
            self._apply_filter()
            self._save_excel(silent=True)
            self._set_status(f"{imported}개 항목 가져오기 완료")
        except Exception as e:
            messagebox.showerror("오류", str(e))

    # ── CRUD ──────────────────────────────────────────────────────────────

    def _commit_images(self, rec, result):
        """다이얼로그 결과에서 이미지를 파일시스템에 반영하고 rec['이미지']를 갱신."""
        existing  = result["_existing"]
        pending   = result["_pending"]
        removed   = result["_removed"]

        # 삭제 처리
        for fname in removed:
            p = full_image_path(rec["ID"], fname)
            if os.path.exists(p):
                os.remove(p)

        # 신규 복사
        for src in pending:
            fname = save_image(src, rec["ID"])
            existing.append(fname)

        rec["이미지"] = ",".join(existing)

    def _add(self):
        dlg = EntryDialog(self.root, title="항목 추가")
        self.root.wait_window(dlg)
        if dlg.result is None:
            return
        rec = {
            "ID":   str(self.next_id),
            "날짜": dlg.result["날짜"],
            "카테고리": dlg.result["카테고리"],
            "태그":     dlg.result["태그"],
            "제목":     dlg.result["제목"],
            "내용":     dlg.result["내용"],
            "이미지":   "",
        }
        self.next_id += 1
        self.records.append(rec)
        self._commit_images(rec, dlg.result)
        self._refresh_categories()
        self._apply_filter()
        self._save_excel(silent=True)
        self._select_by_id(rec["ID"])
        self._set_status("항목 추가됨")

    def _edit(self):
        rec = self._selected_record()
        if rec is None:
            messagebox.showinfo("알림", "수정할 항목을 선택해 주세요.")
            return
        dlg = EntryDialog(self.root, title="항목 수정", data=rec, record_id=rec["ID"])
        self.root.wait_window(dlg)
        if dlg.result is None:
            return
        rec.update({
            "날짜":     dlg.result["날짜"],
            "카테고리": dlg.result["카테고리"],
            "태그":     dlg.result["태그"],
            "제목":     dlg.result["제목"],
            "내용":     dlg.result["내용"],
        })
        self._commit_images(rec, dlg.result)
        self._refresh_categories()
        self._apply_filter()
        self._save_excel(silent=True)
        self._select_by_id(rec["ID"])
        self._set_status("항목 수정됨")

    def _delete(self):
        rec = self._selected_record()
        if rec is None:
            messagebox.showinfo("알림", "삭제할 항목을 선택해 주세요.")
            return
        if not messagebox.askyesno("삭제 확인", f"'{rec.get('제목', '')}' 항목을 삭제할까요?"):
            return
        # 이미지 폴더 삭제
        img_folder = os.path.join(IMG_DIR, rec["ID"])
        if os.path.isdir(img_folder):
            shutil.rmtree(img_folder, ignore_errors=True)
        self.records.remove(rec)
        self._refresh_categories()
        self._apply_filter()
        self._save_excel(silent=True)
        self._show_detail(None)
        self._set_status("항목 삭제됨")

    # ── 필터 & 표시 ───────────────────────────────────────────────────────

    def _refresh_categories(self):
        cats = sorted({r.get("카테고리", "").strip()
                       for r in self.records if r.get("카테고리", "").strip()})
        self.cat_list.delete(0, "end")
        self.cat_list.insert("end", "전체")
        for c in cats:
            self.cat_list.insert("end", c)
        self.cat_list.selection_set(0)

    def _selected_category(self):
        sel = self.cat_list.curselection()
        if not sel:
            return None
        v = self.cat_list.get(sel[0])
        return None if v == "전체" else v

    def _apply_filter(self):
        cat   = self._selected_category()
        query = self.search_var.get().strip().lower()

        def match(r):
            if cat and r.get("카테고리", "").strip() != cat:
                return False
            if query:
                hay = " ".join([r.get("제목",""), r.get("내용",""),
                                r.get("태그",""), r.get("카테고리","")]).lower()
                return query in hay
            return True

        self.filtered = [r for r in self.records if match(r)]
        self._render_tree()
        self._set_status(f"{len(self.filtered)}개 항목")

    def _render_tree(self):
        self.tree.delete(*self.tree.get_children())
        for rec in self.filtered:
            has_img = "📎" if images_for(rec) else ""
            self.tree.insert("", "end", iid=rec["ID"],
                             values=(rec.get("날짜",""), rec.get("카테고리",""),
                                     rec.get("태그",""), rec.get("제목",""),
                                     _content_preview(rec.get("내용","")), has_img))

    def _on_select(self, _=None):
        self._show_detail(self._selected_record())

    def _show_detail(self, rec):
        self.detail_text.configure(state="normal")
        content_load(self.detail_text, rec.get("내용", "") if rec else "")
        self.detail_text.configure(state="disabled")
        self._show_images(rec)

    def _selected_record(self):
        sel = self.tree.selection()
        if not sel:
            return None
        rid = sel[0]
        return next((r for r in self.records if r["ID"] == rid), None)

    def _select_by_id(self, rid):
        if self.tree.exists(rid):
            self.tree.selection_set(rid)
            self.tree.see(rid)
            self._show_detail(self._selected_record())

    def _sort(self, col):
        self.filtered.sort(key=lambda r: r.get(col, ""))
        self._render_tree()

    def _set_status(self, msg):
        self.status_var.set(msg)


def main():
    root = tk.Tk()
    app = WorkNotesApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
